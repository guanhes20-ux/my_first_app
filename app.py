import os, json, sqlite3, secrets, math, io, time, hmac
from pathlib import Path
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from functools import wraps
from collections import defaultdict, deque
from flask import Flask, request, jsonify, render_template, session, redirect, abort, send_file
from cryptography.fernet import Fernet
from docx import Document
from openpyxl import Workbook

BASE = Path(__file__).parent
TZ = ZoneInfo("Europe/Kaliningrad")
CONSENT_VERSION = "2026-10-27-v1"
CARS = json.loads((BASE / "cars.json").read_text())
CAR_MAP = {c["id"]: c for c in CARS}

def now():
    return datetime.now(TZ)

def create_app(config=None):
    app = Flask(__name__)
    app.config.update(SECRET_KEY=os.environ.get("APP_SECRET"), DATA_KEY=os.environ.get("DATA_KEY"),
        ADMIN_PASSWORD=os.environ.get("ADMIN_PASSWORD", ""), PUBLIC_URL=os.environ.get("PUBLIC_URL", "http://localhost:8000").rstrip("/"),
        DATA_DIR=os.environ.get("DATA_DIR", str(BASE / "data")), DEMO_MODE=os.environ.get("DEMO_MODE", "true").lower() == "true",
        BOT_USERNAME=os.environ.get("BOT_USERNAME", "Renta_Car_777_bot"),
        SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", MAX_CONTENT_LENGTH=32768,
        SESSION_COOKIE_SECURE=os.environ.get("PUBLIC_URL", "").startswith("https://"))
    if config:
        app.config.update(config)
    if not app.config["SECRET_KEY"] or len(app.config["SECRET_KEY"]) < 32:
        raise RuntimeError("APP_SECRET must contain at least 32 characters")
    if len(app.config["ADMIN_PASSWORD"]) < 16:
        raise RuntimeError("ADMIN_PASSWORD must contain at least 16 characters")
    cipher = Fernet(app.config["DATA_KEY"])
    if not app.config["DEMO_MODE"]:
        if os.environ.get("LEGAL_REVIEWED") != "true" or not os.environ.get("OPERATOR_NAME") or not os.environ.get("OPERATOR_DETAILS") or not app.config["PUBLIC_URL"].startswith("https://"):
            raise RuntimeError("Production requires HTTPS, operator details and completed legal review")
    folder = Path(app.config["DATA_DIR"])
    folder.mkdir(parents=True, exist_ok=True, mode=0o700)
    db_path = folder / "rental.sqlite3"
    def db():
        conn = sqlite3.connect(db_path, timeout=20)
        conn.row_factory = sqlite3.Row
        return conn
    with db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS bookings (
            id TEXT PRIMARY KEY, car_id TEXT NOT NULL, start TEXT NOT NULL, end TEXT NOT NULL,
            days INTEGER NOT NULL, total INTEGER NOT NULL, status TEXT NOT NULL,
            created TEXT NOT NULL, expires TEXT NOT NULL, telegram_id INTEGER,
            form_token TEXT, form_expires TEXT, personal BLOB, consent_at TEXT, consent_version TEXT)""")
        conn.execute("CREATE TABLE IF NOT EXISTS bot_state (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
    os.chmod(db_path, 0o600)
    app.extensions["rental_db"] = db
    app.extensions["rental_cipher"] = cipher
    limits = defaultdict(deque)

    def limited(bucket, count=30):
        key = (bucket, request.remote_addr)
        entries = limits[key]
        t = time.monotonic()
        while entries and entries[0] < t - 600:
            entries.popleft()
        if len(entries) >= count:
            abort(429)
        entries.append(t)
        # Bound memory for long-running single-worker demo.
        if len(limits) > 10000:
            limits.clear()

    def csrf():
        if not hmac.compare_digest(str(session.get("csrf", "")), str(request.form.get("csrf", "-"))):
            abort(400, "Обновите страницу и повторите действие.")

    def admin(fn):
        @wraps(fn)
        def wrapped(*args, **kwargs):
            if not session.get("admin"):
                return redirect("/admin/login")
            return fn(*args, **kwargs)
        return wrapped

    @app.before_request
    def prepare():
        session.setdefault("csrf", secrets.token_urlsafe(24))

    @app.after_request
    def headers(response):
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["Content-Security-Policy"] = "default-src 'self'; img-src 'self' data:; style-src 'self'; script-src 'self'; form-action 'self'; frame-ancestors 'none'; base-uri 'none'"
        response.headers["Cache-Control"] = "no-store"
        if app.config["SESSION_COOKIE_SECURE"]:
            response.headers["Strict-Transport-Security"] = "max-age=31536000"
        return response

    @app.get("/")
    def index():
        return render_template("index.html", demo=app.config["DEMO_MODE"])

    @app.get("/api/cars")
    def cars():
        return jsonify([{k:v for k,v in c.items() if k not in ("plate","vin")} for c in CARS])

    @app.post("/api/bookings")
    def book():
        limited("booking", 20)
        # Same-origin JSON plus an unguessable session CSRF token.
        if not hmac.compare_digest(request.headers.get("X-CSRF-Token", ""), session["csrf"]):
            abort(400)
        data = request.get_json(silent=True) or {}
        try:
            car = CAR_MAP[data["car_id"]]
            start, end = datetime.fromisoformat(data["start"]), datetime.fromisoformat(data["end"])
            if start.tzinfo is None or end.tzinfo is None:
                raise ValueError()
            start, end = start.astimezone(TZ), end.astimezone(TZ)
            if start < now() + timedelta(minutes=30) or end <= start or end-start > timedelta(days=90):
                raise ValueError()
        except (KeyError, ValueError, TypeError):
            return jsonify(error="Выберите автомобиль и будущий период от 1 часа до 90 суток."), 400
        if end-start < timedelta(hours=1):
            return jsonify(error="Минимальный период — 1 час."), 400
        days = math.ceil((end-start).total_seconds()/86400)
        ident = secrets.token_urlsafe(24)
        with db() as conn:
            conn.execute("BEGIN IMMEDIATE")
            expire_holds(conn)
            conflict = conn.execute("SELECT id FROM bookings WHERE car_id=? AND status IN ('hold','submitted','confirmed') AND start < ? AND end > ?", (car["id"], end.isoformat(), start.isoformat())).fetchone()
            if conflict:
                return jsonify(error="На эти даты автомобиль уже занят. Выберите другие даты или автомобиль."), 409
            conn.execute("INSERT INTO bookings (id,car_id,start,end,days,total,status,created,expires) VALUES (?,?,?,?,?,?,?,?,?)",
                (ident,car["id"],start.isoformat(),end.isoformat(),days,days*car["price"],"hold",now().isoformat(),(now()+timedelta(minutes=30)).isoformat()))
        return jsonify(telegram_url=f"https://t.me/{app.config['BOT_USERNAME']}?start={ident}", days=days, total=days*car["price"])

    def expire_holds(conn):
        conn.execute("UPDATE bookings SET status='expired',form_token=NULL WHERE status='hold' AND expires < ?", (now().isoformat(),))

    @app.route("/application/<token>", methods=["GET","POST"])
    def application(token):
        limited("form", 120)
        with db() as conn:
            expire_holds(conn)
            booking = conn.execute("SELECT * FROM bookings WHERE form_token=? AND status='hold' AND form_expires > ?", (token,now().isoformat())).fetchone()
        if not booking:
            abort(404, "Ссылка истекла или анкета уже отправлена. Начните оформление заново.")
        error = None
        if request.method == "POST":
            csrf()
            personal = {key:request.form.get(key,"").strip() for key in ("name","birth","phone","passport","passport_by","passport_date","address","license","license_date","license_until","experience")}
            import re
            try:
                if not request.form.get("consent") or not request.form.get("terms") or (app.config["DEMO_MODE"] and not request.form.get("fake")):
                    raise ValueError("Подтвердите согласия и использование тестовых данных.")
                if any(not v or len(v)>500 for v in personal.values()):
                    raise ValueError("Заполните все поля (не более 500 символов).")
                if len(personal["name"].split()) < 2 or not re.fullmatch(r"\+7\d{10}",personal["phone"]):
                    raise ValueError("Укажите ФИО и телефон в формате +79991234567.")
                if not re.fullmatch(r"\d{10}",personal["passport"]) or not re.fullmatch(r"\d{10}",personal["license"]):
                    raise ValueError("Серия и номер паспорта и прав должны содержать 10 цифр.")
                birth = datetime.strptime(personal["birth"],"%Y-%m-%d").date()
                start_date = datetime.fromisoformat(booking["start"]).date()
                age = start_date.year-birth.year-((start_date.month,start_date.day)<(birth.month,birth.day))
                if age < 21 or age > 100:
                    raise ValueError("Условия учебного проката: возраст от 21 года.")
                issued = datetime.strptime(personal["license_date"],"%Y-%m-%d").date()
                expiry = datetime.strptime(personal["license_until"],"%Y-%m-%d").date()
                passport_date = datetime.strptime(personal["passport_date"],"%Y-%m-%d").date()
                if issued < birth or issued > now().date() or passport_date < birth or passport_date > now().date() or expiry < datetime.fromisoformat(booking["end"]).date():
                    raise ValueError("Проверьте даты выдачи документов и срок действия прав.")
                exp = int(personal["experience"])
                if exp < 2 or exp > age-16:
                    raise ValueError("Стаж должен быть от 2 лет и соответствовать возрасту.")
                with db() as conn:
                    conn.execute("BEGIN IMMEDIATE")
                    expire_holds(conn)
                    result = conn.execute("UPDATE bookings SET personal=?,consent_at=?,consent_version=?,status='submitted',form_token=NULL WHERE id=? AND status='hold' AND expires > ?",
                        (cipher.encrypt(json.dumps(personal,ensure_ascii=False).encode()),now().isoformat(),CONSENT_VERSION,booking["id"],now().isoformat()))
                    if result.rowcount != 1:
                        abort(409, "Время оформления истекло.")
                return render_template("message.html", title="Анкета отправлена", message="Менеджер проверит заявку и документы. Заявка сама по себе не является подписанным договором. Вернитесь в бот и отправьте /status.")
            except (ValueError, OverflowError) as exc:
                error = str(exc) if str(exc) else "Проверьте введённые данные."
        return render_template("application.html", booking=booking,car=CAR_MAP[booking["car_id"]],demo=app.config["DEMO_MODE"],error=error)

    @app.get("/privacy")
    def privacy():
        return render_template("privacy.html",demo=app.config["DEMO_MODE"],operator=os.environ.get("OPERATOR_NAME","Учебный проект — оператор не задан"),details=os.environ.get("OPERATOR_DETAILS","Реквизиты будут заполнены перед реальным запуском."),version=CONSENT_VERSION)

    @app.route("/admin/login",methods=["GET","POST"])
    def login():
        error = None
        if request.method == "POST":
            limited("login", 10)
            csrf()
            if hmac.compare_digest(request.form.get("password",""),app.config["ADMIN_PASSWORD"]):
                session.clear()
                session["admin"] = True
                session["csrf"] = secrets.token_urlsafe(24)
                session.permanent = True
                app.permanent_session_lifetime = timedelta(hours=2)
                return redirect("/admin")
            error = "Неверный пароль"
        return render_template("login.html",error=error)

    @app.post("/admin/logout")
    @admin
    def logout():
        csrf()
        session.clear()
        return redirect("/")

    @app.get("/admin")
    @admin
    def dashboard():
        with db() as conn:
            expire_holds(conn)
            rows = conn.execute("SELECT * FROM bookings ORDER BY created DESC").fetchall()
        records = [decorate(r) for r in rows]
        return render_template("admin.html",records=records,total=len(records),clients=len({r["telegram_id"] for r in records if r["personal"]}),
            revenue=sum(r["total"] for r in records if r["status"]=="confirmed"))

    def decorate(row):
        result = dict(row)
        result["car"] = CAR_MAP[row["car_id"]]
        result["personal"] = json.loads(cipher.decrypt(row["personal"]).decode()) if row["personal"] else {}
        return result

    @app.post("/admin/<ident>/<action>")
    @admin
    def action(ident,action):
        csrf()
        with db() as conn:
            if action == "confirm":
                conn.execute("UPDATE bookings SET status='confirmed' WHERE id=? AND status='submitted'",(ident,))
            elif action == "cancel":
                conn.execute("UPDATE bookings SET status='cancelled',form_token=NULL WHERE id=? AND status IN ('hold','submitted','confirmed')",(ident,))
            elif action == "delete":
                conn.execute("DELETE FROM bookings WHERE id=?",(ident,))
            else:
                abort(404)
        return redirect("/admin")

    @app.get("/admin/report.xlsx")
    @admin
    def report():
        with db() as conn:
            rows = [decorate(r) for r in conn.execute("SELECT * FROM bookings ORDER BY start")]
        wb = Workbook()
        ws = wb.active
        ws.title = "Заявки"
        ws.append(["ID","Клиент","Телефон","Автомобиль","Получение (Калининград)","Возврат (Калининград)","Сутки","Сумма, ₽","Статус","Создана"])
        for r in rows:
            # Force literal text; prevent formula injection in spreadsheets.
            values = [r["id"],r["personal"].get("name",""),r["personal"].get("phone",""),r["car"]["name"],r["start"],r["end"],r["days"],r["total"],r["status"],r["created"]]
            ws.append(values)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):
                    cell.data_type = "s"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 26
        summary = wb.create_sheet("Статистика")
        summary.append(["Показатель","Значение"])
        summary.append(["Всего заявок",len(rows)])
        summary.append(["Уникальных клиентов",len({r["telegram_id"] for r in rows if r["personal"]})])
        summary.append(["Подтверждено",sum(r["status"]=="confirmed" for r in rows)])
        summary.append(["Сумма подтвержденных заявок, ₽",sum(r["total"] for r in rows if r["status"]=="confirmed")])
        summary.append(["Это сумма договоров, а не учёт фактических оплат",""])
        for c in CARS:
            summary.append([c["name"],sum(r["car_id"]==c["id"] for r in rows)])
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return send_file(stream,as_attachment=True,download_name="rental-report.xlsx")

    @app.get("/admin/<ident>/contract.docx")
    @admin
    def contract(ident):
        with db() as conn:
            row = conn.execute("SELECT * FROM bookings WHERE id=? AND status='confirmed'",(ident,)).fetchone()
        if not row:
            abort(404)
        r = decorate(row)
        p = r["personal"]
        doc = Document()
        for copy in (1,2):
            if copy == 2:
                doc.add_page_break()
            doc.add_heading("ДОГОВОР АРЕНДЫ ТРАНСПОРТНОГО СРЕДСТВА БЕЗ ЭКИПАЖА",0)
            doc.add_paragraph(f"Экземпляр {copy} из 2. № {ident[:10]}. г. Калининград")
            doc.add_paragraph("УЧЕБНЫЙ ШАБЛОН. Требуется проверка реквизитов, условий и документов сторонами перед подписанием.")
            doc.add_paragraph(f"Арендодатель: {os.environ.get('OPERATOR_NAME','[ФИО / наименование]')}. {os.environ.get('OPERATOR_DETAILS','[Адрес, ИНН, ОГРН/ОГРНИП, банковские реквизиты]')}")
            doc.add_paragraph(f"Арендатор: {p['name']}, дата рождения {p['birth']}; паспорт {p['passport']}, выдан {p['passport_by']} {p['passport_date']}; адрес: {p['address']}; телефон: {p['phone']}. ВУ {p['license']}, выдано {p['license_date']}, действительно до {p['license_until']}.")
            doc.add_heading("1. Предмет и срок",1)
            doc.add_paragraph(f"Арендодатель передаёт во временное владение и пользование без экипажа автомобиль {r['car']['name']}, {r['car']['year']} г., госномер {r['car']['plate']}, VIN {r['car']['vin']}. Срок: {r['start']} — {r['end']} (UTC+02:00). Передача и возврат оформляются актами.")
            doc.add_heading("2. Плата и обеспечительный платёж",1)
            doc.add_paragraph(f"Тариф {r['car']['price']} ₽ за начатые 24 часа. {r['days']} суток; сумма {r['total']} ₽. Обеспечительный платёж {r['car']['deposit']} ₽ не входит в арендную плату. Порядок оплаты, основания удержания и срок возврата: [согласовать и заполнить].")
            doc.add_heading("3. Права и обязанности",1)
            doc.add_paragraph("Арендодатель передаёт исправный автомобиль с документами и действующим ОСАГО с допустимым кругом водителей. Арендатор использует автомобиль по назначению, соблюдает ПДД, возвращает его в согласованном состоянии с учётом нормального износа. Условия обслуживания, ремонта, разрешённая территория, лимит пробега и топливо: [согласовать].")
            doc.add_heading("4. Ответственность и прекращение",1)
            doc.add_paragraph("Ответственность определяется законодательством РФ, в том числе § 3 главы 34 ГК РФ. Повреждения, расходы и причинная связь документируются. Порядок действий при ДТП, страхование, досрочное прекращение, продление, просрочка возврата и урегулирование споров: [согласовать]. Условия, ограничивающие обязательные права потребителя, не применяются.")
            doc.add_heading("5. Подписи и акт передачи / возврата",1)
            doc.add_paragraph("Договор составлен в двух экземплярах равной силы, по одному каждой стороне. Отправка анкеты и подтверждение менеджером не заменяют подписание.")
            doc.add_paragraph("Передача: дата ______ пробег ______ топливо ______ комплектность ______ повреждения ______\nВозврат: дата ______ пробег ______ топливо ______ комплектность ______ повреждения ______\nАрендодатель __________ /__________/  Арендатор __________ /__________/")
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return send_file(stream,as_attachment=True,download_name=f"contract-{ident[:10]}-2-copies.docx")

    @app.cli.command("purge-expired")
    def purge_expired():
        cutoff = (now()-timedelta(days=int(os.environ.get("RETENTION_DAYS","30")))).isoformat()
        with db() as conn:
            expire_holds(conn)
            result = conn.execute("DELETE FROM bookings WHERE (status IN ('expired','cancelled') AND created < ?) OR (status='hold' AND expires < ?)",(cutoff,cutoff))
        print(f"Deleted {result.rowcount} expired/cancelled records. Review submitted/confirmed records separately.")
    return app

if __name__ == "__main__":
    create_app().run(host="127.0.0.1",port=8000,debug=False)
