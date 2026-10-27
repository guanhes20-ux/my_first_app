import io
from datetime import timedelta
import pytest
from cryptography.fernet import Fernet
from docx import Document
from openpyxl import load_workbook
from app import create_app, now
from bot import handle

@pytest.fixture()
def app(tmp_path):
    return create_app({"TESTING":True,"SECRET_KEY":"x"*40,"DATA_KEY":Fernet.generate_key(),
        "ADMIN_PASSWORD":"test-password-123456789","DATA_DIR":str(tmp_path),"DEMO_MODE":True,
        "PUBLIC_URL":"https://example.test"})

def create_booking(client, car="rio", shift=0):
    client.get("/")
    with client.session_transaction() as s:
        csrf=s["csrf"]
    start=(now()+timedelta(days=2+shift)).replace(microsecond=0)
    return client.post("/api/bookings",json={"car_id":car,"start":start.isoformat(),
        "end":(start+timedelta(hours=25)).isoformat()},headers={"X-CSRF-Token":csrf})

def booking_id(response):
    return response.json["telegram_url"].split("start=")[1]

def claim(app, ident, uid=123):
    messages=[]
    handle(app,{"from":{"id":uid},"chat":{"id":uid,"type":"private"},"text":"/start "+ident},
        lambda *args:messages.append(args))
    return messages

def form_csrf(client,path):
    assert client.get(path).status_code==200
    with client.session_transaction() as s:
        return s["csrf"]

def personal():
    return {"name":"Тестовый Клиент","birth":"1990-01-01","phone":"+79991234567",
        "passport":"0000000000","passport_by":"Учебный орган","passport_date":"2010-01-01",
        "address":"Вымышленный адрес","license":"0000000000","license_date":"2020-01-01",
        "license_until":"2040-01-01","experience":"5","consent":"on","terms":"on","fake":"on"}

def login(client):
    csrf=form_csrf(client,"/admin/login")
    return client.post("/admin/login",data={"csrf":csrf,"password":"test-password-123456789"})

def test_price_overlap_and_adjacent(app):
    client=app.test_client()
    r=create_booking(client)
    assert r.status_code==200 and r.json["days"]==2 and r.json["total"]==4600
    assert create_booking(client).status_code==409
    assert create_booking(client,shift=3).status_code==200
    assert create_booking(client,car="k5").status_code==200

def test_expired_hold_releases_car(app):
    client=app.test_client()
    ident=booking_id(create_booking(client))
    with app.extensions["rental_db"]() as db:
        db.execute("UPDATE bookings SET expires=? WHERE id=?",((now()-timedelta(minutes=1)).isoformat(),ident))
    assert create_booking(client).status_code==200
    assert "недоступна" in claim(app,ident)[0][1]

def test_csrf_and_dates(app):
    client=app.test_client()
    client.get("/")
    assert client.post("/api/bookings",json={}).status_code==400
    with client.session_transaction() as s:
        csrf=s["csrf"]
    assert client.post("/api/bookings",json={"car_id":"rio","start":"2020-01-01","end":"2020-01-02"},headers={"X-CSRF-Token":csrf}).status_code==400
    assert client.get("/admin/report.xlsx").status_code==302
    assert client.get("/application/unknown").status_code==404

def test_claim_protection_and_group_rejection(app):
    client=app.test_client()
    ident=booking_id(create_booking(client))
    assert claim(app,ident)[0][2]["inline_keyboard"]
    assert len(claim(app,ident,uid=456)[0])==2
    out=[]
    handle(app,{"from":{"id":123},"chat":{"id":123,"type":"group"},"text":"/start "+ident},lambda *args:out.append(args))
    assert out==[]

def test_end_to_end_documents_encryption_and_delete(app):
    client=app.test_client()
    ident=booking_id(create_booking(client))
    msg=claim(app,ident)
    path=msg[0][2]["inline_keyboard"][0][0]["url"].replace("https://example.test","")
    data=personal()
    data["csrf"]=form_csrf(client,path)
    assert client.post(path,data=data).status_code==200
    assert client.get(path).status_code==404
    with app.extensions["rental_db"]() as db:
        row=db.execute("SELECT * FROM bookings WHERE id=?",(ident,)).fetchone()
    assert row["status"]=="submitted"
    assert "Тестовый".encode() not in row["personal"]
    assert login(client).status_code==302
    with client.session_transaction() as s:
        csrf=s["csrf"]
    assert client.post("/admin/"+ident+"/confirm",data={"csrf":csrf}).status_code==302
    result=client.get("/admin/"+ident+"/contract.docx")
    assert result.status_code==200
    doc=Document(io.BytesIO(result.data))
    text="\n".join(p.text for p in doc.paragraphs)
    assert "Экземпляр 1 из 2" in text and "Экземпляр 2 из 2" in text
    assert text.count("Тестовый Клиент")==2
    report=load_workbook(io.BytesIO(client.get("/admin/report.xlsx").data))
    assert report["Заявки"]["B2"].value=="Тестовый Клиент"
    assert report["Статистика"]["B5"].value==4600
    assert client.post("/admin/"+ident+"/delete",data={"csrf":csrf}).status_code==302
    with app.extensions["rental_db"]() as db:
        assert db.execute("SELECT count(*) FROM bookings").fetchone()[0]==0

def test_missing_consent_rejected(app):
    client=app.test_client()
    ident=booking_id(create_booking(client))
    path=claim(app,ident)[0][2]["inline_keyboard"][0][0]["url"].replace("https://example.test","")
    data=personal()
    del data["consent"]
    data["csrf"]=form_csrf(client,path)
    response=client.post(path,data=data)
    assert "Подтвердите".encode() in response.data
    with app.extensions["rental_db"]() as db:
        assert db.execute("SELECT status FROM bookings WHERE id=?",(ident,)).fetchone()[0]=="hold"

def test_bot_delete_is_scoped(app):
    client=app.test_client()
    first=booking_id(create_booking(client))
    second=booking_id(create_booking(client,car="k5"))
    claim(app,first,123)
    claim(app,second,456)
    handle(app,{"from":{"id":123},"chat":{"id":123,"type":"private"},"text":"/delete"},lambda *a:None)
    with app.extensions["rental_db"]() as db:
        rows=db.execute("SELECT id FROM bookings").fetchall()
    assert [r["id"] for r in rows]==[second]
